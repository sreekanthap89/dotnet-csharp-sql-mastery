from pathlib import Path
from typing import Dict, Any
import pypdf
import docx

def parse_file(file_path: Path) -> Dict[str, Any]:
    """
    Parses a supported file format (.md, .txt, .pdf, .docx) and extracts plain text and metadata.
    """
    suffix = file_path.suffix.lower()
    filename = file_path.name
    
    if suffix in [".md", ".markdown", ".txt", ".json", ".cs", ".sql"]:
        with open(file_path, "r", encoding="utf-8", errors="replace") as f:
            content = f.read()
        return {
            "title": filename,
            "filename": filename,
            "file_type": suffix.lstrip("."),
            "text": content,
            "size_bytes": file_path.stat().st_size
        }
        
    elif suffix == ".pdf":
        text_parts = []
        reader = pypdf.PdfReader(str(file_path))
        for page_idx, page in enumerate(reader.pages):
            page_text = page.extract_text() or ""
            if page_text.strip():
                text_parts.append(f"--- Page {page_idx + 1} ---\n{page_text}")
        return {
            "title": filename,
            "filename": filename,
            "file_type": "pdf",
            "text": "\n\n".join(text_parts),
            "page_count": len(reader.pages),
            "size_bytes": file_path.stat().st_size
        }
        
    elif suffix in [".docx", ".doc"]:
        doc = docx.Document(str(file_path))
        text_parts = []
        for p in doc.paragraphs:
            if p.text.strip():
                text_parts.append(p.text)
        for table in doc.tables:
            for row in table.rows:
                row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
                if row_text:
                    text_parts.append(row_text)
        return {
            "title": filename,
            "filename": filename,
            "file_type": "docx",
            "text": "\n\n".join(text_parts),
            "size_bytes": file_path.stat().st_size
        }
        
    else:
        # Fallback text reading
        try:
            with open(file_path, "r", encoding="utf-8", errors="replace") as f:
                content = f.read()
            return {
                "title": filename,
                "filename": filename,
                "file_type": suffix.lstrip(".") or "unknown",
                "text": content,
                "size_bytes": file_path.stat().st_size
            }
        except Exception as e:
            raise ValueError(f"Unsupported or unreadable file format: {suffix} ({e})")
